"""
medicine_matcher.py — The brain of the prescription reader.

Loads your medicine database (the verbatim-names xlsx the doctor verified) and
matches OCR'd text lines against it using fuzzy matching. This is what turns a
garbled OCR read like "Eitlax Md 0.25" into a confident hit on "ETILAX MD 0.25".

Key ideas
---------
1. The DB is the single source of truth. This module READS it; it does not contain
   any medicine data itself. Edit the Excel, and every run uses the new version.

2. Matching is fuzzy (rapidfuzz), not exact. OCR is never character-perfect, so we
   score similarity and only auto-accept above a threshold. Below it -> flagged for
   the human. A MISSING match is safe (gets reviewed); a WRONG match is dangerous,
   so the threshold is deliberately strict.

3. The ALIASES column is the learning memory. Every confirmed correction of a
   misread gets appended there, so the same misread maps instantly next time.
   (Sarvam itself remembers nothing between calls — this local DB is the memory.)

DB columns expected (case-insensitive, extra columns are fine):
    - "Medicine name (verbatim)"   the bill string  -> the canonical match target
    - "Salt / Composition ..."     filled by you/doctor
    - "Category"                    optional
    - "Use case (doctor)"           optional
    - "Conf"                        optional (Verified / Confident / FLAG)
    - "Aliases"                     optional; ';'-separated known misreads (the memory)

If your column headers differ, set them in COLUMN_MAP below.
"""

import os
import re
import openpyxl
from rapidfuzz import process, fuzz

# ---- Tunables -------------------------------------------------------------
AUTO_ACCEPT_SCORE = 88   # >= this: confident match, auto-filled
REVIEW_FLOOR      = 70   # between floor and accept: show as "likely" but flag
                         # < floor: treated as NO match (flag, offer top guesses)
TOP_N_CANDIDATES  = 3    # how many guesses to show for a weak/no match
# ---------------------------------------------------------------------------

# Map logical fields -> possible header spellings in your xlsx (first match wins).
COLUMN_MAP = {
    "name":     ["medicine name (verbatim)", "medicine name (verbatim from bill)",
                 "medicine name", "name", "row labels"],
    "salt":     ["salt / composition (search-derived)", "salt / composition",
                 "verified salt", "salt", "composition"],
    "category": ["category"],
    "use":      ["use case (doctor)", "use case", "indication", "use"],
    "conf":     ["conf", "confidence"],
    "aliases":  ["aliases", "aliases (raw billing variants)", "alias"],
}


def _norm(s):
    """Loose normaliser used to widen matching: lowercase, strip punctuation,
    collapse whitespace. Keeps digits (strengths help disambiguate 0.25 vs 0.5)."""
    s = (s or "").lower()
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


class MedicineDB:
    def __init__(self, xlsx_path):
        self.path = xlsx_path
        self.records = []          # list of dicts (one per DB row)
        self._match_index = {}     # normalised_string -> record index
        self._headers = {}
        self._header_row = None
        self._name_col = None
        self._alias_col = None
        self.load()

    # ---------- loading ----------
    def _resolve_columns(self, ws):
        # find the header row (the row containing the 'name' header)
        for r in range(1, min(ws.max_row, 12) + 1):
            rowvals = [(_norm(c.value) if c.value else "") for c in ws[r]]
            for cand in COLUMN_MAP["name"]:
                if _norm(cand) in rowvals:
                    self._header_row = r
                    break
            if self._header_row:
                break
        if not self._header_row:
            raise ValueError("Could not find a header row containing a medicine-name column.")

        headers = {}
        for idx, cell in enumerate(ws[self._header_row], start=1):
            h = _norm(cell.value) if cell.value else ""
            headers[idx] = h
        # map logical field -> column index
        col_for = {}
        for field, options in COLUMN_MAP.items():
            for opt in options:
                no = _norm(opt)
                for idx, h in headers.items():
                    if h == no:
                        col_for[field] = idx
                        break
                if field in col_for:
                    break
        if "name" not in col_for:
            raise ValueError("Medicine-name column not found after header detection.")
        self._headers = col_for
        self._name_col = col_for["name"]
        self._alias_col = col_for.get("aliases")

    def load(self):
        wb = openpyxl.load_workbook(self.path, data_only=True)
        ws = wb.active
        self._resolve_columns(ws)
        self.records = []
        self._match_index = {}
        for r in range(self._header_row + 1, ws.max_row + 1):
            name = ws.cell(r, self._name_col).value
            if not name or str(name).strip().upper() == "TOTAL":
                continue
            name = str(name).strip()
            rec = {"row": r, "name": name}
            for field, idx in self._headers.items():
                if field == "name":
                    continue
                val = ws.cell(r, idx).value
                rec[field] = (str(val).strip() if val is not None else "")
            # build match keys: the name itself + every alias
            keys = [name]
            aliases = rec.get("aliases", "")
            if aliases:
                keys += [a.strip() for a in aliases.split(";") if a.strip()]
            rec["_keys"] = keys
            ridx = len(self.records)
            self.records.append(rec)
            for k in keys:
                self._match_index[_norm(k)] = ridx
        # the flat list of normalised strings rapidfuzz searches over
        self._choices = list(self._match_index.keys())

    # ---------- matching ----------
    def match_line(self, raw_text):
        """Match one OCR text fragment. Returns a dict with the decision."""
        q = _norm(raw_text)
        if not q:
            return {"raw": raw_text, "status": "empty"}

        # exact normalised hit (covers names + learned aliases) -> instant 100
        if q in self._match_index:
            rec = self.records[self._match_index[q]]
            return self._hit(raw_text, rec, 100.0, "exact")

        results = process.extract(q, self._choices, scorer=fuzz.WRatio,
                                  limit=TOP_N_CANDIDATES)
        if not results:
            return {"raw": raw_text, "status": "no_match", "candidates": []}

        best_key, best_score, _ = results[0]
        best_rec = self.records[self._match_index[best_key]]

        cands = [{"name": self.records[self._match_index[k]]["name"],
                  "salt": self.records[self._match_index[k]].get("salt", ""),
                  "score": round(sc, 1)} for k, sc, _ in results]

        if best_score >= AUTO_ACCEPT_SCORE:
            return self._hit(raw_text, best_rec, best_score, "fuzzy", cands)
        elif best_score >= REVIEW_FLOOR:
            d = self._hit(raw_text, best_rec, best_score, "weak", cands)
            d["status"] = "review"      # likely, but human confirms
            return d
        else:
            return {"raw": raw_text, "status": "no_match",
                    "best_score": round(best_score, 1), "candidates": cands}

    def _hit(self, raw, rec, score, how, cands=None):
        return {
            "raw": raw,
            "status": "matched" if score >= AUTO_ACCEPT_SCORE else "review",
            "match": rec["name"],
            "salt": rec.get("salt", ""),
            "category": rec.get("category", ""),
            "use": rec.get("use", ""),
            "db_conf": rec.get("conf", ""),
            "score": round(score, 1),
            "how": how,
            "candidates": cands or [],
        }

    # ---------- learning (alias memory) ----------
    def add_alias(self, correct_name, misread_text):
        """Append a confirmed misread as an alias of the correct medicine, IN the
        Excel, so future runs match it instantly. Returns True on success.

        NOTE: call this only for APPROVED corrections — a wrong alias permanently
        teaches a wrong mapping. (See learn_corrections.py for the approval flow.)"""
        misread = misread_text.strip()
        if not misread:
            return False
        wb = openpyxl.load_workbook(self.path)   # not data_only: we write back
        ws = wb.active
        # ensure an aliases column exists
        if self._alias_col is None:
            new_col = max(self._headers.values()) + 1
            ws.cell(self._header_row, new_col, "Aliases")
            self._alias_col = new_col
        # find the row for correct_name
        target_row = None
        for rec in self.records:
            if _norm(rec["name"]) == _norm(correct_name):
                target_row = rec["row"]
                break
        if target_row is None:
            return False
        cur = ws.cell(target_row, self._alias_col).value or ""
        existing = {_norm(a) for a in cur.split(";") if a.strip()}
        if _norm(misread) in existing:
            return True   # already known
        new_val = (cur + "; " if cur.strip() else "") + misread
        ws.cell(target_row, self._alias_col, new_val)
        wb.save(self.path)
        self.load()   # refresh in-memory index so it's live immediately
        return True


if __name__ == "__main__":
    # quick self-test against the salts-filled DB
    import sys
    db_path = sys.argv[1] if len(sys.argv) > 1 else "All_Medicines_Salts_Filled.xlsx"
    db = MedicineDB(db_path)
    print(f"Loaded {len(db.records)} medicines, {len(db._choices)} match keys.\n")
    for test in ["Eitlax Md 0.25", "Q-pin SR 100", "valkem od 500",
                 "Dynaper AQ", "totally unknown drug xyz"]:
        r = db.match_line(test)
        print(f"{test!r:28s} -> {r['status']:9s} "
              f"{r.get('match','-')}  (score {r.get('score','-')})")
